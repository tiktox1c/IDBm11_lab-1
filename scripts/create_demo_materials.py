"""Создать два учебных источника PDF/DOCX. Вопросов и ответов теста здесь нет."""
from html import escape
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from ai_test_generator.config import OUTPUT, TEMP, prepare_temp


def create():
    import pymupdf
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor

    prepare_temp()
    folder = OUTPUT / 'demo_materials'
    folder.mkdir(exist_ok=True)
    for number, name in enumerate(('hardware', 'learning'), 1):
        source = (ROOT / 'examples' / (name+'.md')).read_text(encoding='utf-8-sig')
        sections = source.strip().split('\n---\n')
        document = Document()
        document.sections[0].page_width = Mm(210)
        document.sections[0].page_height = Mm(297)
        for side in ('top_margin','bottom_margin','left_margin','right_margin'):
            setattr(document.sections[0],side,Mm(19))
        for style in ('Normal','Title','Heading 1'):
            document.styles[style].font.name = 'Arial'
            document.styles[style].font.color.rgb = RGBColor(0,0,0)
        document.styles['Normal'].font.size = Pt(11)
        document.styles['Normal'].paragraph_format.space_after = Pt(6)
        document.core_properties.author = 'AI Test Generator'
        pdf = pymupdf.open()
        expected = []
        for page_index, section in enumerate(sections):
            if page_index:
                document.add_page_break()
            html = []
            for block in re.split(r'\n\s*\n',section.strip()):
                if block.startswith('# '):
                    text,tag = block[2:], 'h1'
                    document.add_paragraph(text,style='Title')
                elif block.startswith('## '):
                    text,tag = block[3:], 'h2'
                    document.add_paragraph(text,style='Heading 1')
                else:
                    text,tag = block,'p'
                    document.add_paragraph(text)
                expected.append(text)
                html.append(f'<{tag}>'+escape(text).replace('\n','<br>')+f'</{tag}>')
            page = pdf.new_page(width=595,height=842)
            spare, scale = page.insert_htmlbox(pymupdf.Rect(54,48,541,794), ''.join(html),
                css='body{font-family:sans-serif;font-size:11pt;line-height:1.35;color:#111}'
                    'h1{font-size:18pt;margin:0 0 14pt}h2{font-size:14pt;margin:12pt 0 8pt}'
                    'p{margin:0 0 9pt}',scale_low=1)
            if spare < 0:
                raise ValueError('Учебный материал не помещается на страницу; нужен перенос.')
        stem = 'Учебный материал-'+str(number)
        pdf_path,docx_path = folder/(stem+'.pdf'),folder/(stem+'.docx')
        pdf.save(pdf_path)
        pdf.close()
        document.save(docx_path)
        # Проверяем фактическое наличие всех слов после чтения обоих форматов.
        from ai_test_generator.extraction import extract
        compact = lambda text: re.sub(r'\s+','',text)
        for path in (pdf_path, docx_path):
            material=extract(path)
            actual=compact(material.text)
            if not all(compact(block) in actual for block in expected):
                raise ValueError(f'Часть исходного текста недоступна после сохранения {path.name}')
            print(path)
        with pymupdf.open(pdf_path) as pdf:
            for index,page in enumerate(pdf):
                page.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(TEMP/f'demo-{number}-page-{index+1}.png')


if __name__ == '__main__':
    create()
