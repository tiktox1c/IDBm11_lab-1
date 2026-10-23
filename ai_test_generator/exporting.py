import json
import os
import tempfile
import errno
import shutil
from pathlib import Path

from .config import TEMP, OUTPUT, TYPES, prepare_temp
from .validation import validate_test


def answer_text(q):
    if q["type"] == "choice":
        return f"{q['answer']}. {q['options'][q['answer'] - 1]}"
    if q["type"] == "sequence":
        return " → ".join(map(str, q["answer"]))
    if q["type"] == "matching":
        return "; ".join(f"{i}–{v}" for i, v in enumerate(q["answer"], 1))
    return q["answer"]


def export_test(test, path=None):
    path = Path(path or OUTPUT / "test.docx").expanduser().resolve()
    source = test.get("source_path")
    if source:
        source = Path(source).resolve()
        if path == source or (path.exists() and source.exists() and path.samefile(source)):
            raise ValueError("Нельзя перезаписывать исходный документ. Выберите другое имя.")
    if path.suffix.lower() not in (".docx", ".json"):
        raise ValueError("Выберите формат DOCX или JSON.")
    validate_test(test)
    prepare_temp()
    # Сначала полностью записываем во временный файл, затем заменяем целевой.
    fd, name = tempfile.mkstemp(suffix=path.suffix, dir=TEMP)
    os.close(fd)
    temporary = Path(name)
    try:
        if path.suffix.lower() == ".json":
            temporary.write_text(json.dumps(test, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            from docx import Document
            from docx.shared import Pt, Cm
            from .review import question_status
            from .attachments import export_attachments, format_table
            doc = Document()
            section=doc.sections[0]
            section.page_width=Cm(21);section.page_height=Cm(29.7)
            section.top_margin=section.bottom_margin=Cm(2)
            section.left_margin=section.right_margin=Cm(2)
            doc.styles["Normal"].font.name = "Arial"
            doc.styles["Normal"].font.size = Pt(11)
            doc.add_heading(test.get('title') or "AI Test Generator — учебный тест", 0)
            doc.add_paragraph("Дисциплина: Программно-аппаратное обеспечение искусственного интеллекта")
            doc.add_paragraph(f"Источник: {test['source_name']}\nМодель: {test['model']}")
            doc.add_paragraph(f"Вопросов: {len(test['questions'])} из запрошенных {test.get('requested_count', 8)}. Требуется проверка преподавателем.")
            for i, q in enumerate(test["questions"], 1):
                question_start=len(doc.paragraphs)
                doc.add_heading(f"{i}. {TYPES[q['type']]}", 2)
                doc.add_paragraph(q["prompt"])
                doc.add_paragraph("Статус: " + question_status(q))
                export_attachments(doc,q.get("attachments",[]))
                if q["type"] in ("choice", "sequence"):
                    key = "options" if q["type"] == "choice" else "items"
                    for j, item in enumerate(q[key], 1):
                        doc.add_paragraph(f"{j}. {item}")
                    if q["type"] == "sequence":
                        doc.add_paragraph("Ответ: номера элементов в правильном порядке: __________________")
                elif q["type"] == "matching":
                    table = doc.add_table(rows=1, cols=2)
                    table.style = "Table Grid"
                    table.rows[0].cells[0].text = "Левая колонка"
                    table.rows[0].cells[1].text = "Правая колонка"
                    for j, (left, right) in enumerate(zip(q["left"], q["right"]), 1):
                        cells = table.add_row().cells
                        cells[0].text, cells[1].text = f"{j}. {left}", f"{j}. {right}"
                    format_table(table)
                    doc.add_paragraph("Ответ: пары номеров: __________________")
                else:
                    doc.add_paragraph("Ответ: __________________________________________________")
                block=doc.paragraphs[question_start:]
                for paragraph in block[:-1]:paragraph.paragraph_format.keep_with_next=True
                if block:block[-1].paragraph_format.keep_with_next=False
            doc.add_page_break()
            doc.add_heading("Правильные ответы и основания", 1)
            if any(question_status(q)=='Требует проверки' for q in test['questions']):
                doc.add_paragraph('Вопросы с отметкой «ТРЕБУЕТ ПРОВЕРКИ» содержат предложенные моделью ответы; они не подтверждены преподавателем.')
            for i, q in enumerate(test["questions"], 1):
                doc.add_paragraph(f"{i}. {answer_text(q)} — Статус: {question_status(q)}")
                if q.get('acceptable_answers'):
                    doc.add_paragraph('Допустимые формулировки: ' + '; '.join(q['acceptable_answers']))
                if question_status(q)=='Требует проверки':
                    doc.add_paragraph('ТРЕБУЕТ ПРОВЕРКИ: ' + '; '.join(q.get('review_reasons', [])))
                elif q.get('manual_confirmed'):
                    doc.add_paragraph('Проверено вручную пользователем.')
                doc.add_paragraph(f"Источник: {q['source_reference']}\nЦитата: {q['source_quote']}")
            if test.get("warnings"):
                doc.add_heading("Замечания", 1)
                for warning in test["warnings"]:
                    doc.add_paragraph(warning)
            doc.save(temporary)
        try:
            os.replace(temporary, path)
        except OSError as exc:
            if exc.errno != errno.EXDEV and getattr(exc, "winerror", None) != 17:
                raise
            # Временный файл всегда на E:, пользователь вправе выбрать другой диск.
            shutil.copyfile(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
