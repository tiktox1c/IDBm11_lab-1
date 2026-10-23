"""Вложения вопросов: исходные таблицы, изображения и сохранённые OMML."""
from pathlib import Path
from docx.shared import Cm
from docx.oxml import parse_xml, OxmlElement

def export_attachments(doc,items):
    for item in items:
        kind=item['type']
        if kind=='table':
            rows=item.get('rows',[])
            if not rows:continue
            table=doc.add_table(rows=0,cols=max(map(len,rows)));table.style='Table Grid'
            for values in rows:
                cells=table.add_row().cells
                for cell,value in zip(cells,values):cell.text=str(value or '')
            format_table(table,bool(item.get('headers')))
            if item.get('title'):doc.add_paragraph(item['title'])
        elif kind=='formula' and item.get('omml'):
            doc.add_paragraph()._p.append(parse_xml(item['omml']))
        elif item.get('path'):
            path=Path(item['path'])
            if not path.is_file():raise ValueError('Вложение недоступно: '+item.get('reference',path.name))
            try:
                from docx.image.image import Image
                image=Image.from_file(str(path))
                width=min(16,21*image.px_width/image.px_height)
                doc.add_picture(str(path),width=Cm(width))
            except Exception as exc:raise ValueError('Word не поддерживает формат вложения. Прикрепите фрагмент исходной страницы: '+item.get('reference',path.name)) from exc
        elif kind=='formula' and item.get('reliable'):
            doc.add_paragraph(item['text'])
        else:raise ValueError('Вложение не подготовлено. Откройте его в просмотре материала перед сохранением.')
        doc.add_paragraph('Материал: '+item['reference'])


def format_table(table,header=True):
    for i,row in enumerate(table.rows):
        properties=row._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        if i==0 and header:properties.append(OxmlElement('w:tblHeader'))
        if len(table.rows)<=7:
            for cell in row.cells:
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=i<len(table.rows)-1
