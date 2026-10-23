import copy
import hashlib
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.oxml import parse_xml
from ai_test_generator.config import TEMP, prepare_temp
from ai_test_generator.extraction import extract
from ai_test_generator.review import confirm_question, save_correction, return_to_review, question_status
from ai_test_generator.exporting import export_test
from test_core import fixture

class ReleaseTests(unittest.TestCase):
    def setUp(self):
        prepare_temp()
        self.directory = tempfile.TemporaryDirectory(dir=TEMP)
        self.base = Path(self.directory.name)
    def tearDown(self):
        self.directory.cleanup()
    def test_statuses_require_explicit_actions_and_invalid_structure_is_blocked(self):
        q=fixture('short')
        self.assertEqual(question_status(q),'Требует проверки')
        confirm_question(q)
        self.assertEqual(question_status(q),'Проверено')
        changed=copy.deepcopy(q); changed['answer']='Процессор CPU'
        save_correction(q,changed)
        self.assertEqual(question_status(q),'Исправлено')
        return_to_review(q)
        self.assertEqual(question_status(q),'Требует проверки')
        broken=fixture('choice'); broken['answer']=9
        with self.assertRaises(ValueError):confirm_question(broken)
    def test_model_cannot_forge_review_status(self):
        from ai_test_generator.generation import generate_test
        from ai_test_generator.extraction import Material
        def model(prompt,schema):
            return {'question':'Какой вычислитель назван в источнике?', 'correct_answer':'CPU', 'wrong_answers':['мышь','кабель','экран'], 'review_status':'checked', 'manual_confirmed':True}
        result=generate_test(Material('x','Вычислитель CPU. RAM — память; SSD — накопитель; GPU — ускоритель.'),call=model,count=4)
        for q in result['questions']:self.assertEqual(question_status(q),'Требует проверки')
    def test_docx_table_image_formula_roundtrip_and_unchanged_source(self):
        import pymupdf
        image=self.base/'image.png'
        pix=pymupdf.Pixmap(pymupdf.csRGB,(0,0,100,60));pix.clear_with(180);pix.save(image)
        doc=Document();doc.add_paragraph('Учебный материал. Рисунок 1 — локальная схема CPU.')
        doc.add_picture(str(image))
        table=doc.add_table(rows=3,cols=2)
        for row,values in zip(table.rows,[['Компонент','Назначение'],['CPU','Вычисления'],['RAM','Память']]):
            for cell,value in zip(row.cells,values):cell.text=value
        p=doc.add_paragraph('Формула: ')
        p._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'))
        source=self.base/'source.docx';doc.save(source);before=hashlib.sha256(source.read_bytes()).digest()
        material=extract(source)
        self.assertEqual(before,hashlib.sha256(source.read_bytes()).digest())
        types={e['type'] for e in material.elements}
        self.assertTrue({'text','table','image','formula'}<=types)
        self.assertTrue(all(e['source_path']==str(source) and e['reference'] for e in material.elements))
        tab=next(e for e in material.elements if e['type']=='table')
        self.assertEqual(tab['rows'][1],['CPU','Вычисления']);self.assertIn('Компонент',tab['text'])
        formula=next(e for e in material.elements if e['type']=='formula')
        self.assertEqual(formula['text'],'(a)/(b)');self.assertTrue(formula['reliable'])
        picture=next(e for e in material.elements if e['type']=='image')
        self.assertTrue(Path(picture['path']).is_file())
        q=fixture('short');q['attachments']=[tab,formula,picture];confirm_question(q)
        test={'questions':[q],'source_name':source.name,'source_path':str(source),'model':'qwen2.5:1.5b','title':'Проверка'}
        target=self.base/'result.docx';export_test(test,target);result=Document(target)
        self.assertTrue(result.tables);self.assertEqual(len(result.inline_shapes),1)
        self.assertIn('oMath',result.element.xml)
        self.assertIn('Статус: Проверено','\n'.join(p.text for p in result.paragraphs))
        self.assertAlmostEqual(result.sections[0].page_width.cm,21,places=1)
    def test_uncertain_pdf_table_page_is_not_used_for_automatic_answers(self):
        import pymupdf
        from unittest.mock import patch
        from types import SimpleNamespace
        path=self.base/'uncertain.pdf'
        with pymupdf.open() as doc:
            p=doc.new_page();p.insert_text((40,40),'CPU');p.insert_text((250,40),'Unrelated row values')
            doc.save(path)
        table=SimpleNamespace(bbox=(0,0,590,800),cells=[(0,0,590,800)],extract=lambda:[['CPU','Wrong relation']])
        with patch.object(pymupdf.Page,'find_tables',return_value=SimpleNamespace(tables=[table])):material=extract(path)
        self.assertIn('CPU',material.text)
        self.assertNotIn('CPU',material.generation_text)
    def test_merged_docx_table_is_available_but_not_used_for_generation(self):
        doc=Document();doc.add_paragraph('Обычный учебный текст о процессоре.')
        table=doc.add_table(rows=2,cols=2);table.cell(0,0).merge(table.cell(0,1)).text='UNRELIABLE_HEADER'
        table.cell(1,0).text='SECRET_VALUE';table.cell(1,1).text='UNCERTAIN_PAIR'
        path=self.base/'merged.docx';doc.save(path);material=extract(path)
        self.assertIn('SECRET_VALUE',material.text)
        self.assertNotIn('SECRET_VALUE',material.generation_text)

    def test_tall_image_fits_printable_page(self):
        import pymupdf
        image=self.base/'tall.png';pix=pymupdf.Pixmap(pymupdf.csRGB,(0,0,20,800));pix.clear_with(100);pix.save(image)
        q=fixture('short');q['attachments']=[{'type':'image','path':str(image),'reference':'Свой рисунок'}]
        test={'questions':[q],'source_name':'example','model':'qwen2.5:1.5b'}
        target=self.base/'tall.docx';export_test(test,target)
        shape=Document(target).inline_shapes[0]
        self.assertLessEqual(shape.height.cm,22)
        self.assertLessEqual(shape.width.cm,17)

    def test_unknown_omml_is_preserved_but_not_used_as_fact(self):
        doc=Document();doc.add_paragraph('Доступный текст о процессоре CPU.')
        doc.add_paragraph()._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:unknown><m:r><m:t>broken</m:t></m:r></m:unknown></m:oMath>'))
        path=self.base/'formula.docx';doc.save(path);material=extract(path)
        formula=next(e for e in material.elements if e['type']=='formula')
        self.assertFalse(formula['reliable']);self.assertNotIn('broken',material.generation_text)
        self.assertIn('oMath',formula['omml'])
    def test_pdf_table_is_structured_once_images_and_formula_kept(self):
        import pymupdf
        path=self.base/'table.pdf'
        with pymupdf.open() as doc:
            p=doc.new_page();p.insert_text((40,40),'Equation: x = a + b')
            for x in (40,170,300):p.draw_line((x,80),(x,200))
            for y in (80,120,160,200):p.draw_line((40,y),(300,y))
            for y,values in zip((105,145,185),[('Part','Role'),('CPU','Compute'),('RAM','Memory')]):
                p.insert_text((45,y),values[0]);p.insert_text((175,y),values[1])
            pix=pymupdf.Pixmap(pymupdf.csRGB,(0,0,20,20));pix.clear_with(200)
            p.insert_image(pymupdf.Rect(40,230,80,270),pixmap=pix);doc.save(path)
        before=path.read_bytes();material=extract(path)
        self.assertEqual(before,path.read_bytes())
        table=next(e for e in material.elements if e['type']=='table')
        self.assertTrue(table['reliable']);self.assertEqual(material.text.count('Compute'),1)
        self.assertTrue(any(e['type']=='image' and e.get('path') for e in material.elements))
        self.assertTrue(any(e['type']=='formula' and 'x = a + b' in e['text'] for e in material.elements))

if __name__=='__main__':unittest.main()
