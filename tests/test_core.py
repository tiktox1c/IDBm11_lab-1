"""Синтетические данные — только для проверок; приложение их не импортирует."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from urllib.error import URLError

from ai_test_generator.config import MODEL, TEMP, prepare_temp
from ai_test_generator.extraction import Material, extract
from ai_test_generator.exporting import export_test
from ai_test_generator.generation import chunks, generate_test, ollama_generate, OllamaError
from ai_test_generator.validation import validate_question

prepare_temp()


def fixture(kind="choice"):
    q = {"type": kind, "prompt": "Какой вычислитель назван в источнике?",
         "source_quote": "Вычислитель CPU. Этапы: чтение, обработка, запись, вывод. RAM — память; SSD — накопитель; GPU — ускоритель.",
         "source_reference": "synthetic", "edited": False}
    q["source_excerpt"] = q["source_quote"]
    q.update({"choice": {"options": ["CPU", "мышь", "экран", "кабель"], "answer": 1},
              "sequence": {"items": ["запись", "чтение", "обработка"], "answer": [2, 3, 1]},
              "matching": {"left": ["RAM", "SSD", "GPU"], "right": ["ускоритель", "память", "накопитель"], "answer": [2, 3, 1]},
              "short": {"answer": "CPU"}}[kind])
    return q


class ValidationTests(unittest.TestCase):
    def test_four_kinds(self):
        for kind in ("choice", "sequence", "matching", "short"):
            with self.subTest(kind=kind):
                validate_question(fixture(kind))

    def test_bad_answers(self):
        cases = [("choice", 0), ("choice", True), ("choice", 5), ("short", ""),
                 ("sequence", [1, 1, 2]), ("matching", [1, 2]), ("matching", [True, 2, 3])]
        for kind, answer in cases:
            q = fixture(kind)
            q["answer"] = answer
            with self.subTest(kind=kind, answer=answer), self.assertRaises(ValueError):
                validate_question(q)

    def test_variants_and_grounding(self):
        for change in ({"options": ["CPU", "cpu!", "x", "y"]},
                       {"options": ["CPU", "x", "y"]}, {"source_quote": "Выдуманная цитата"},
                       {"options": ["TPU", "x", "y", "z"]}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                validate_question({**fixture(), **change})

    def test_duplicate(self):
        q = fixture()
        with self.assertRaises(ValueError):
            validate_question({**q, "prompt": q["prompt"] + "!"}, [q])

    def test_chunks_cover_text(self):
        text = "\n".join(f"Строка {i}: учебный материал." for i in range(200))
        parts = chunks(text)
        self.assertTrue(all(len(p) <= 1600 for p in parts))
        self.assertTrue(text.endswith(parts[-1]))
        for line in text.splitlines():
            self.assertTrue(any(line in part for part in parts))


class IOTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=TEMP)
        self.base = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def test_docx_text_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Учебный материал об ИИ")
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text, cells[1].text = "CPU", "Процессор"
        path = self.base / "source.docx"
        doc.save(path)
        result = extract(path)
        self.assertIn("Учебный материал", result.text)
        self.assertIn("CPU | Процессор", result.text)

    def test_pdf_text_table_and_empty_page(self):
        import pymupdf
        path = self.base / "source.pdf"
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((50, 50), "AI hardware: CPU and RAM")
            for x in (50, 150, 250):
                page.draw_line((x, 100), (x, 180))
            for y in (100, 140, 180):
                page.draw_line((50, y), (250, y))
            page.insert_text((60, 125), "CPU")
            page.insert_text((160, 125), "Processor")
            page.insert_text((60, 165), "RAM")
            page.insert_text((160, 165), "Memory")
            doc.new_page()
            doc.save(path)
        material = extract(path)
        self.assertIn("AI hardware", material.text)
        self.assertIn("CPU | Processor", material.text)
        self.assertTrue(any("Страница 2" in w for w in material.warnings))

    def test_empty_pdf_rejected(self):
        import pymupdf
        path = self.base / "empty.pdf"
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(path)
        with self.assertRaisesRegex(ValueError, "Доступный текст"):
            extract(path)

    def test_export_roundtrip(self):
        from docx import Document
        questions = [fixture(kind) for kind in ("choice", "sequence", "matching", "short")]
        for i, q in enumerate(questions):
            q["prompt"] = f"{i}: " + ["Выберите вычислитель", "Расположите этапы", "Сопоставьте термины", "Назовите вычислитель"][i]
        test = {"model": MODEL, "source_name": "synthetic", "questions": questions, "warnings": ["Проверить вручную"]}
        export_test(test, self.base / "result.json")
        self.assertEqual(json.loads((self.base / "result.json").read_text(encoding="utf-8")), test)
        export_test(test, self.base / "result.docx")
        doc = Document(self.base / "result.docx")
        content = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Правильные ответы", content)
        self.assertIn("2 → 3 → 1", content)
        self.assertIn("1–2; 2–3; 3–1", content)
        self.assertEqual(len(doc.tables), 1)

class GenerationTests(unittest.TestCase):
    def test_no_fallback_and_bounded_retries(self):
        def refuse(prompt, schema):
            return {"available": False, "reason": "Недостаточно данных", "question": None}
        with patch("builtins.print"):
            result = generate_test(Material("source", "CPU"), call=refuse)
        self.assertEqual(result["questions"], [])
        self.assertEqual(sum("не создано" in w for w in result["warnings"]), 8)

    def test_partial_failure_preserves_result(self):
        calls = 0

        def call(prompt, schema):
            nonlocal calls
            calls += 1
            if calls == 1:
                return {"available": True, "reason": "", "question": fixture()}
            raise OllamaError("Ollama остановлена")

        result = generate_test(Material("source", fixture()["source_excerpt"]), call=call)
        self.assertEqual(len(result["questions"]), 1)
        self.assertIn("Ollama остановлена", result["warnings"])

    def test_eight_questions_two_each(self):
        counters = dict.fromkeys(("choice", "sequence", "matching", "short"), 0)
        prompts = {"choice": ["Выберите вычислитель.", "Какое устройство хранит данные?"],
                   "sequence": ["Укажите порядок операций.", "Как выполняется цикл обработки?"],
                   "matching": ["Сопоставьте оборудование и назначение.", "Найдите соответствия компонентов и функций."],
                   "short": ["Назовите вычислитель.", "Как называется память?"]}

        def call(prompt, schema):
            fields = schema["properties"]
            kind = "choice" if "correct_answer" in fields else "sequence" if "steps" in fields else "matching" if "pairs" in fields else "short"
            slot = counters[kind]
            counters[kind] += 1
            q = fixture(kind)
            q["prompt"] = prompts[kind][min(slot, 1)]
            if slot:
                if kind == "choice":
                    q["options"] = ["SSD", "a", "b", "c"]
                elif kind == "sequence":
                    q["items"] = ["чтение", "обработка", "запись", "вывод"]
                    q["answer"] = [1, 2, 3, 4]
                elif kind == "matching":
                    q["left"] = ["память", "накопитель", "ускоритель"]
                    q["right"] = ["GPU", "RAM", "SSD"]
                else:
                    q["answer"] = "RAM"
            return {"available": True, "reason": "", "question": q}

        result = generate_test(Material("synthetic", fixture()["source_excerpt"]), call=call)
        self.assertEqual(len(result["questions"]), 8)
        for kind in counters:
            self.assertEqual(sum(q["type"] == kind for q in result["questions"]), 2)

    def test_malformed_response_retried_no_more_than_sixteen_times(self):
        calls = []

        def invalid(prompt, schema):
            calls.append(prompt)
            return {"available": True, "reason": "", "question": {"prompt": "Без ответа"}}

        result = generate_test(Material("synthetic", "CPU"), call=invalid)
        self.assertEqual(len(calls), 16)
        self.assertEqual(result["questions"], [])
        self.assertNotEqual(calls[0], calls[1])

    def test_http_payload_only_fixed_model(self):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                pass

            def read(self):
                return json.dumps({"done": True, "response": '{"available": false}'}).encode()

        with patch("ai_test_generator.generation.build_opener") as opener:
            opener.return_value.open.return_value = Response()
            ollama_generate("Материал", {"type": "object"})
            request = opener.return_value.open.call_args.args[0]
            payload = json.loads(request.data)
            self.assertEqual(payload["model"], "qwen2.5:1.5b")
            self.assertEqual(payload["options"]["num_gpu"], 0)
            self.assertEqual(request.full_url, "http://127.0.0.1:11434/api/generate")

    def test_connection_failure_is_explained(self):
        with patch("ai_test_generator.generation.build_opener") as opener:
            opener.return_value.open.side_effect = URLError("refused")
            with self.assertRaises(OllamaError):
                ollama_generate("text", {})

    def test_invalid_json_details_stay_out_of_user_warnings(self):
        def invalid(prompt, schema):
            return json.loads("{not valid JSON}")

        result = generate_test(Material("synthetic", "CPU"), call=invalid)
        self.assertEqual(result["questions"], [])
        self.assertFalse(any("line 1" in warning or "column" in warning for warning in result["warnings"]))


if __name__ == "__main__":
    unittest.main()
