"""Extraction regressions: preserve text order, provenance and original files."""
import hashlib
from collections import Counter
from pathlib import Path
import re
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from ai_test_generator.extraction import extract, read_pdf, table_text


class ExtractionRevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def pdf_table(self, *, merged=False):
        import pymupdf
        path = self.folder / 'table.pdf'
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((50, 50), 'Section before table')
            for y in (90, 130, 170):
                page.draw_line((50, y), (300, y))
            for x in (50, 300):
                page.draw_line((x, 90), (x, 170))
            page.draw_line((175, 130 if merged else 90), (175, 170))
            page.insert_text((60, 115), 'Combined heading' if merged else 'Component')
            if not merged:
                page.insert_text((185, 115), 'Function')
            page.insert_text((60, 155), 'CPU')
            page.insert_text((185, 155), 'Processor')
            page.insert_text((50, 210), 'Section after table')
            doc.new_page()
            doc.save(path)
        return path

    def test_pdf_table_cells_occur_once_in_reading_order(self):
        path = self.pdf_table()
        original = path.read_bytes()
        material = extract(path)
        self.assertEqual(material.text.count('CPU'), 1)
        self.assertEqual(material.text.count('Processor'), 1)
        self.assertIn('CPU | Processor', material.text)
        self.assertLess(material.text.index('Section before'), material.text.index('[Таблица'))
        self.assertLess(material.text.index('[Таблица'), material.text.index('Section after'))
        self.assertIn('[Страница 2]', material.text)
        self.assertEqual(path.read_bytes(), original)

    def test_pdf_merged_cells_are_not_duplicated(self):
        material = extract(self.pdf_table(merged=True))
        self.assertEqual(material.text.count('Combined heading'), 1)
        self.assertIn('CPU | Processor', material.text)

    def test_filled_paragraph_backgrounds_are_not_tables(self):
        import pymupdf
        path = self.folder / 'backgrounds.pdf'
        with pymupdf.open() as doc:
            page = doc.new_page()
            for y, line in ((100, 'Background paragraph first'), (120, 'Background paragraph second'),
                            (140, 'Background paragraph third')):
                page.draw_rect((50, y - 15, 300, y + 5), color=None, fill=(1, 1, 1))
                page.insert_text((60, y), line)
            doc.save(path)
        text = extract(path).text
        self.assertEqual(text.count('Background paragraph'), 3)
        self.assertNotIn('[Таблица', text)

    def test_pdf_table_detection_failure_preserves_accessible_text(self):
        import pymupdf
        path = self.pdf_table()
        with patch.object(pymupdf.Page, 'find_tables', side_effect=RuntimeError('fixture failure')):
            material = extract(path)
        self.assertIn('Section before table', material.text)
        self.assertIn('CPU', material.text)
        self.assertTrue(any('таблицы не удалось' in w for w in material.warnings))

    def test_pdf_incomplete_large_table_detection_does_not_erase_prose(self):
        import pymupdf
        path = self.pdf_table()
        # Detector geometry can span a paragraph although extracted cells omit
        # some text. Reject the replacement and retain the original text layer.
        table = SimpleNamespace(bbox=(0, 0, 600, 800), cells=[(0, 0, 600, 800)],
                                extract=lambda: [['CPU', 'Processor']])
        with patch.object(pymupdf.Page, 'find_tables', return_value=SimpleNamespace(tables=[table])):
            text = extract(path).text
        self.assertIn('Section before table', text)
        self.assertIn('Section after table', text)
        self.assertEqual(text.count('CPU'), 1)
        self.assertNotIn('[Таблица', text)

    def test_blank_pdf_rejected_even_with_page_markers(self):
        import pymupdf
        path = self.folder / 'blank.pdf'
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(path)
        with self.assertRaisesRegex(ValueError, 'Доступный текст'):
            extract(path)

    def test_corrupt_pdf_reports_readable_error(self):
        path = self.folder / 'corrupt.pdf'
        path.write_bytes(b'not a PDF')
        with self.assertRaisesRegex(ValueError, 'PDF'):
            extract(path)

    def test_password_pdf_reports_protection(self):
        import pymupdf
        path = self.folder / 'protected.pdf'
        with pymupdf.open() as doc:
            doc.new_page().insert_text((50, 50), 'Secret text')
            doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
        with self.assertRaisesRegex(ValueError, 'паролем'):
            extract(path)

    def test_docx_merged_cells_and_nested_tables_appear_once_in_order(self):
        from docx import Document
        doc = Document()
        doc.add_heading('Hardware section', level=1)
        doc.add_paragraph('First listed fact', style='List Bullet')
        table = doc.add_table(rows=2, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = 'Merged heading'
        cell = table.cell(1, 0)
        cell.text = 'Before nested'
        cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Nested detail'
        cell.add_paragraph('After nested')
        table.cell(1, 1).text = 'Other cell'
        doc.add_paragraph('Final paragraph')
        path = self.folder / 'merged.docx'
        doc.save(path)
        original = hashlib.sha256(path.read_bytes()).digest()
        text = extract(path).text
        self.assertEqual(text.count('Merged heading'), 1)
        self.assertEqual(text.count('Nested detail'), 1)
        positions = [text.index(s) for s in ('Hardware section', 'First listed fact', 'Merged heading',
                    'Before nested', 'Nested detail', 'After nested', 'Final paragraph')]
        self.assertEqual(positions, sorted(positions))
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), original)

    def test_docx_vertical_merge_does_not_repeat_cell(self):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0)).text = 'Vertical heading'
        table.cell(0, 1).text = 'Upper detail'
        table.cell(1, 1).text = 'Lower detail'
        path = self.folder / 'vertical.docx'
        doc.save(path)
        text = extract(path).text
        self.assertEqual(text.count('Vertical heading'), 1)
        self.assertIn('Upper detail', text)
        self.assertIn('Lower detail', text)

    def test_docx_automatic_list_labels_keep_order_and_bullets_distinct(self):
        from docx import Document
        doc = Document()
        doc.add_heading('Process', level=1)
        doc.add_paragraph('Read input', style='List Number')
        doc.add_paragraph('Calculate result', style='List Number')
        doc.add_paragraph('Write output', style='List Number')
        doc.add_paragraph('Processor', style='List Bullet')
        doc.add_paragraph('Memory', style='List Bullet')
        path = self.folder / 'lists.docx'
        doc.save(path)
        text = extract(path).text
        for item in ('1. Read input', '2. Calculate result', '3. Write output',
                     '• Processor', '• Memory'):
            self.assertIn(item, text)

    def test_docx_list_hierarchy_and_explicit_restart_are_retained(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()

        def element(name, **attributes):
            node = OxmlElement('w:' + name)
            for key, value in attributes.items():
                node.set(qn('w:' + key), str(value))
            return node

        numbering = doc.part.numbering_part.element
        abstract = element('abstractNum', abstractNumId=90)
        for index, start, pattern in ((0, 3, '%1.'), (1, 1, '%1.%2.')):
            level = element('lvl', ilvl=index)
            level.extend([element('start', val=start), element('numFmt', val='decimal'),
                          element('lvlText', val=pattern)])
            abstract.append(level)
        numbering.append(abstract)
        for number_id in (90, 91):
            num = element('num', numId=number_id)
            num.append(element('abstractNumId', val=90))
            if number_id == 91:
                override = element('lvlOverride', ilvl=0)
                override.append(element('startOverride', val=7))
                num.append(override)
            numbering.append(num)
        for text, num_id, level in (('Gather', 90, 0), ('Read', 90, 1),
                                    ('Compute', 90, 1), ('Finish', 90, 0), ('Restart', 91, 0)):
            paragraph = doc.add_paragraph(text)
            num = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num.get_or_add_numId().val = num_id
            num.get_or_add_ilvl().val = level
        path = self.folder / 'hierarchy.docx'
        doc.save(path)
        text = extract(path).text
        for expected in ('3. Gather', '3.1. Read', '3.2. Compute', '4. Finish', '7. Restart'):
            self.assertIn(expected, text)

    def test_table_text_keeps_zero_and_false_values(self):
        self.assertEqual(table_text([[0, False, None]]), '0 | False | ')

    def test_available_real_pdf_documents_keep_all_text_once_and_inputs_unchanged(self):
        import pymupdf
        paths = sorted((Path(__file__).resolve().parents[1] / 'text').glob('*.pdf'))
        if not paths:
            self.skipTest('The optional local PDF corpus is not present')
        for path in paths:
            with self.subTest(document=path.name):
                original = hashlib.sha256(path.read_bytes()).digest()
                with pymupdf.open(path) as doc:
                    source = ''.join(page.get_text(sort=True) for page in doc)
                    pages = len(doc)
                material = extract(path)
                content = re.sub(r'\[Страница \d+\]', '', material.text)
                self.assertEqual(Counter(c for c in content if not c.isspace()),
                                 Counter(c for c in source if not c.isspace()))
                for number in range(1, pages + 1):
                    self.assertIn(f'[Страница {number}]', material.text)
                self.assertEqual(material.name, path.name)
                self.assertEqual(material.source_path, str(path.resolve()))
                self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), original)


if __name__ == '__main__':
    unittest.main()
